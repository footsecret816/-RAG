from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from common import (
    normalize_material,
    normalize_size_system,
    parse_number,
    product_data_root,
    split_values,
    text,
    write_json,
)


COLUMN_ALIASES = {
    "SKU_ID": ["SKU_ID", "SKU", "货号", "编号", "产品编号"],
    "Product_Category": ["Product_Category", "产品分类", "品类"],
    "Factory_Name": ["Factory_Name", "Factory", "工厂", "生产厂家", "供应商"],
    "MOQ": ["MOQ", "起订量", "最小起订量"],
    "Price": ["Price", "价格", "工厂价", "单价"],
    "Material": ["Material", "材质", "主材质"],
    "Material_Detail": ["Material_Detail", "详细材质", "材质明细", "产品材质"],
    "Size_System": ["Size_System", "尺码体系", "尺码制"],
    "Size_Range": ["Size_Range", "尺码", "尺码范围"],
    "Description": ["Description", "描述", "产品特性", "特性", "功能/描述", "功能描述"],
    "Function_Tags": ["Function_Tags", "功能标签"],
    "Scenario_Tags": ["Scenario_Tags", "场景标签", "使用场景"],
    "Special_Features": ["Special_Features", "特殊特性", "特殊属性"],
    "Cushioning": ["Cushioning", "缓震性", "缓震"],
    "Elasticity": ["Elasticity", "回弹性", "回弹"],
    "Softness": ["Softness", "软硬度", "软硬"],
    "Arch_Height": ["Arch_Height", "足弓高度"],
    "Arch_Support": ["Arch_Support", "支撑强度", "足弓支撑强度"],
    "Heel_Cup_Depth": ["Heel_Cup_Depth", "后跟杯深度"],
}


def _rows_from_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _rows_from_xlsx(path: Path, sheet: str | None) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 xlsx 需要 openpyxl：python -m pip install openpyxl") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    iterator = ws.iter_rows(values_only=True)
    try:
        headers = [text(v) for v in next(iterator)]
    except StopIteration:
        return []

    rows: list[dict[str, Any]] = []
    for values in iterator:
        row = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        if any(not (v is None or text(v) == "") for v in row.values()):
            rows.append(row)
    return rows


def read_rows(path: Path, sheet: str | None = None) -> list[dict[str, Any]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return _rows_from_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _rows_from_xlsx(path, sheet)
    raise ValueError("V1 仅支持 .xlsx/.xlsm/.csv 固定格式表格")


def map_headers(row: dict[str, Any]) -> dict[str, Any]:
    normalized_headers = {text(k).casefold(): k for k in row.keys()}
    result: dict[str, Any] = {}
    for target, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            original = normalized_headers.get(alias.casefold())
            if original is not None:
                result[target] = row.get(original)
                break
    return result


def find_image(image_dir: Path | None, sku: str) -> tuple[str | None, list[str]]:
    if image_dir is None or not image_dir.exists():
        return None, []
    exact: list[Path] = []
    prefix: list[Path] = []
    allowed = {".jpg", ".jpeg", ".png", ".webp"}
    for path in image_dir.iterdir():
        if not path.is_file() or path.suffix.casefold() not in allowed:
            continue
        stem = path.stem.casefold()
        if stem == sku.casefold():
            exact.append(path)
        elif stem.startswith(sku.casefold() + "-"):
            prefix.append(path)
    candidates = exact or prefix
    if len(candidates) == 1:
        return str(candidates[0].resolve()), []
    if len(candidates) > 1:
        return None, [str(p.resolve()) for p in candidates]
    return None, []


def build_candidate(
    source: Path,
    rows: list[dict[str, Any]],
    default_category: str,
    image_dir: Path | None,
) -> dict[str, Any]:
    grouped: dict[str, dict[str, Any]] = {}
    issues: list[dict[str, Any]] = []

    for row_no, raw in enumerate(rows, start=2):
        row = map_headers(raw)
        sku = text(row.get("SKU_ID"))
        factory = text(row.get("Factory_Name"))
        if not sku:
            issues.append({"row": row_no, "type": "blocking", "message": "缺少 SKU_ID"})
            continue
        if not factory:
            issues.append({"row": row_no, "sku": sku, "type": "blocking", "message": "缺少 Factory_Name"})
            continue

        category = text(row.get("Product_Category")) or default_category
        material_raw = row.get("Material")
        material = normalize_material(material_raw)
        if text(material_raw) and material is None:
            issues.append({
                "row": row_no,
                "sku": sku,
                "factory": factory,
                "type": "review",
                "field": "Material",
                "source_value": text(material_raw),
                "message": "材质无法映射到标准词",
            })

        size_raw = row.get("Size_System")
        size_system = normalize_size_system(size_raw)
        if text(size_raw) and size_system is None:
            issues.append({
                "row": row_no,
                "sku": sku,
                "factory": factory,
                "type": "review",
                "field": "Size_System",
                "source_value": text(size_raw),
                "message": "尺码体系无法映射到 EU/US/UK/Alpha",
            })

        item = grouped.setdefault(
            sku,
            {
                "Product_Category": category,
                "SKU_ID": sku,
                "Main_Image_Source": None,
                "Packaging_Options": [],
                "Function_Tags": [],
                "Scenario_Tags": [],
                "Special_Features": [],
                "Performance_Attributes": {
                    "Cushioning": None,
                    "Elasticity": None,
                    "Softness": None,
                    "Arch_Height": None,
                    "Arch_Support": None,
                    "Heel_Cup_Depth": None,
                },
                "Factory_Offers": [],
                "Source_Descriptions": [],
                "Review_Items": [],
            },
        )

        description = text(row.get("Description"))
        if description and description not in item["Source_Descriptions"]:
            item["Source_Descriptions"].append(description)

        for key in ("Function_Tags", "Scenario_Tags", "Special_Features"):
            for value in split_values(row.get(key)):
                if value not in item[key]:
                    item[key].append(value)

        for key in ("Cushioning", "Elasticity", "Softness"):
            value = parse_number(row.get(key))
            if value is not None:
                current = item["Performance_Attributes"].get(key)
                if current is None:
                    item["Performance_Attributes"][key] = value
                elif current != value:
                    issues.append({
                        "row": row_no, "sku": sku, "type": "review", "field": key,
                        "message": f"同一 SKU 出现冲突值：{current} vs {value}",
                    })

        for key in ("Arch_Height", "Arch_Support", "Heel_Cup_Depth"):
            value = text(row.get(key))
            if value:
                current = item["Performance_Attributes"].get(key)
                if current is None:
                    item["Performance_Attributes"][key] = value
                elif current != value:
                    issues.append({
                        "row": row_no, "sku": sku, "type": "review", "field": key,
                        "message": f"同一 SKU 出现冲突值：{current} vs {value}",
                    })

        offer = {
            "Factory_Name": factory,
            "Material": material,
            "Material_Detail": text(row.get("Material_Detail")) or None,
            "Price": parse_number(row.get("Price")),
            "MOQ": parse_number(row.get("MOQ")),
            "Size_System": size_system,
            "Size_Range": text(row.get("Size_Range")) or None,
        }

        existing = next(
            (x for x in item["Factory_Offers"] if x.get("Factory_Name") == factory),
            None,
        )
        if existing is None:
            item["Factory_Offers"].append(offer)
        else:
            for field, value in offer.items():
                if field == "Factory_Name" or value is None or value == "":
                    continue
                if existing.get(field) in (None, ""):
                    existing[field] = value
                elif existing.get(field) != value:
                    issues.append({
                        "row": row_no,
                        "sku": sku,
                        "factory": factory,
                        "type": "review",
                        "field": field,
                        "message": f"同一 SKU + Factory 出现冲突：{existing.get(field)} vs {value}",
                    })

    for sku, item in grouped.items():
        image, image_candidates = find_image(image_dir, sku)
        item["Main_Image_Source"] = image
        if image_candidates:
            issues.append({
                "sku": sku,
                "type": "review",
                "field": "Main_Image",
                "message": "匹配到多张可能的主图，需要确认",
                "candidates": image_candidates,
            })
        if item["Source_Descriptions"]:
            item["Review_Items"].append({
                "id": f"{sku}-SEMANTIC",
                "status": "pending",
                "field": "Semantic_Fields",
                "source_text": " | ".join(item["Source_Descriptions"]),
                "message": "请按 02 规则由 Agent 生成/确认 Function_Tags、Scenario_Tags、Performance 候选；不得仅凭材质猜测。",
            })

    return {
        "candidate_schema_version": "v1.0",
        "source_file": str(source.resolve()),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "products": list(grouped.values()),
        "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="固定格式产品表格 → 03 候选数据")
    parser.add_argument("source")
    parser.add_argument("--sheet")
    parser.add_argument("--category", default="鞋垫")
    parser.add_argument("--image-dir")
    parser.add_argument("--output")
    args = parser.parse_args()

    source = Path(args.source).expanduser().resolve()
    image_dir = Path(args.image_dir).expanduser().resolve() if args.image_dir else source.parent
    rows = read_rows(source, args.sheet)
    candidate = build_candidate(source, rows, args.category, image_dir)

    root = product_data_root()
    output = (
        Path(args.output).expanduser().resolve()
        if args.output
        else root / "_Staging" / f"candidate-{source.stem}.json"
    )
    write_json(output, candidate)

    blocking = [x for x in candidate["issues"] if x.get("type") == "blocking"]
    print(json.dumps({
        "ok": not blocking,
        "output": str(output),
        "row_count": len(rows),
        "product_count": len(candidate["products"]),
        "blocking_issue_count": len(blocking),
        "review_issue_count": len(candidate["issues"]) - len(blocking),
        "pending_semantic_review_count": sum(
            1 for p in candidate["products"] for x in p.get("Review_Items", [])
            if x.get("status") == "pending"
        ),
    }, ensure_ascii=False, indent=2))
    return 0 if not blocking else 2


if __name__ == "__main__":
    raise SystemExit(main())
